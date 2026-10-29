from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.chart import BarChart, Reference

from emplyees_data import load_employees, find_contact_issues

BASE_DIR = Path(__file__).resolve().parent

TEMPLATE_FILE = BASE_DIR / 'templates' / 'hr_report.xlsx'
OUTPUT_FILE = BASE_DIR / 'output' / 'HR_отчет.xlsx'

START_ROW = 9


def create_excel_report(employees: list[dict], issues: list[dict]) -> Path:
    """..."""

    if not TEMPLATE_FILE.is_file():
        raise FileNotFoundError(f"file not found at {TEMPLATE_FILE}")

    workbook = load_workbook(TEMPLATE_FILE)

    try:
        sheet = workbook["Сотрудники"]

        invalid_codes = {issue["code"] for issue in issues}

        for row_number, employee in enumerate(employees, start=START_ROW):
            status = "Ошибка" if employee["employee_code"] in invalid_codes else "ОК"

            values = (
                employee["id"],
                employee["employee_code"],
                employee["full_name"],
                employee["department"],
                employee["position"],
                employee["salary"],
                employee["email"],
                employee["phone"],
                employee["hire_date"],
                status,
            )

            for column_number, value in enumerate(values, start=1):
                sheet.cell(row=row_number, column=column_number, value=value)

        OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(OUTPUT_FILE)

    finally:
        workbook.close()

    return OUTPUT_FILE


if __name__ == '__main__':
    employees = load_employees()
    issues = find_contact_issues(employees)

    result = create_excel_report(employees, issues)


    print(f'Сотрудников выгружено: {len(employees)}')
    print(f'Ошибок обнаружено: {len(issues)}')
    print(f'Файл сохранен: {result}')




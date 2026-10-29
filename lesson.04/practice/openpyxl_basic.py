from pathlib import Path
from openpyxl import load_workbook, Workbook


BASE_DIR = Path(__file__).resolve().parents[1]

OUTPUT_FILE = BASE_DIR / 'output' / 'temp_excel.xlsx'


workbook = Workbook()
sheet = workbook.active
sheet.title = "Сотрудники"
print(f"Листы книги: {workbook.sheetnames}")


sheet["A1"] = "Код"
sheet["B1"] = "ФИО"
sheet["C1"] = "ОКЛАД"

print(f'Заголовок A1: {sheet["A1"].value}')


sheet.cell(row=2, column=1, value="ЕМР-0001")
sheet.cell(row=2, column=2, value="Волков Александр Иванович")
sheet.cell(row=2, column=3, value="115700")

print(f'Код сотрудника: {sheet["A2"].value}')
print(f'Оклад сотрудника: {sheet.cell(row=2, column=3).value}')

print()
print()


sheet.append(["EMP-0002", "Макарова Ксения Дмитриевна", 97000])
sheet.append(["EMP-0003", "Соловьёв Тимофей Михайлович", 125400])
sheet.append(["EMP-0004", "Новикова Ольга Алексеевна", 35400])

print(f'Третий сотрудник: {sheet["B4"].value}')
print(f'Последняя строка: {sheet.max_row}')

print()
print()

notes_sheet = workbook.create_sheet("Примечания")
notes_sheet["A1"] = "Тестовая книга"

employees_sheet = workbook["Сотрудники"]

print(f"Все Листы книги: {workbook.sheetnames}")


##########################

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

workbook.save(OUTPUT_FILE)
workbook.close()